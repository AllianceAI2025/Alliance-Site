"""Render the shared demo content as an editable firm-template Word draft."""
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def color(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for attribute in list(fonts.attrib):
        if "theme" in attribute.lower():
            del fonts.attrib[attribute]
    fonts.set(qn("w:eastAsia"), name)


def cell_format(cell, fill=None):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn("w:" + key), value)
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement("w:tcMar")
    for side, size in (("top", 110), ("bottom", 110), ("left", 115), ("right", 115)):
        part = OxmlElement("w:" + side)
        part.set(qn("w:w"), str(size))
        part.set(qn("w:type"), "dxa")
        margins.append(part)
    props.append(margins)
    if fill:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        props.append(shading)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def write(brief, output):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.12)
    section.bottom_margin = Inches(.68)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance = Inches(.38)
    section.footer_distance = Inches(.32)
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False

    styles = doc.styles
    color(styles["Normal"], "Calibri", 11.5)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    styles["Normal"].paragraph_format.space_after = Pt(7)
    color(styles["Title"], "Georgia", 25)
    for border in styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    styles["Title"].paragraph_format.space_before = Pt(7)
    styles["Title"].paragraph_format.space_after = Pt(6)
    for heading in ("Heading 1", "Heading 2"):
        color(styles[heading], "Calibri", 12, True)
        styles[heading].paragraph_format.space_before = Pt(12)
        styles[heading].paragraph_format.space_after = Pt(5)
        styles[heading].paragraph_format.keep_with_next = True
    color(styles["Heading 1"], "Calibri", 18, True)
    for name, face, size in (("Firm Letterhead", "Georgia", 22), ("Firm Practice", "Calibri", 8.5), ("Matter Metadata", "Calibri", 10), ("Source Reference", "Calibri", 9)):
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        color(style, face, size)
        style.paragraph_format.space_after = Pt(3)
    styles["Source Reference"].paragraph_format.line_spacing = 1.05

    # Populate every header/footer variant to preserve the firm's letterhead
    # even when Word applies a user's first-page or facing-page settings.
    for part in (section.header, section.first_page_header, section.even_page_header):
        header = part.paragraphs[0]
        header.clear()
        header.style = styles["Firm Letterhead"]
        header.text = brief["firm"]
        header.paragraph_format.space_after = Pt(2)
        practice = part.add_paragraph(brief["practice"].upper(), "Firm Practice")
        practice.paragraph_format.space_after = Pt(0)
    for part in (section.footer, section.first_page_footer, section.even_page_footer):
        footer = part.paragraphs[0]
        footer.clear()
        footer.style = styles["Source Reference"]
        footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.25))
        footer.add_run("Fictional firm and matter · Internal draft")
        footer.add_run("\t")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

    doc.add_paragraph(brief["name"], "Matter Metadata")
    doc.add_paragraph(brief["documentTitle"], "Title")
    sub = doc.add_paragraph(brief["subtitle"])
    sub.paragraph_format.space_after = Pt(13)
    for label, value in (("To", brief["reviewer"]), ("From", brief["author"]), ("Matter", brief["matter"] + " · " + brief["client"]), ("Date", brief["date"] + " · Revision " + str(brief["revision"]) + " · " + brief["status"])):
        p = doc.add_paragraph(style="Matter Metadata")
        p.add_run(label + "   ").bold = True
        p.add_run(value)
    doc.add_paragraph(brief["introduction"])
    for part in brief["sections"]:
        doc.add_heading(part["title"], 2)
        doc.add_paragraph(part["body"])

    doc.add_page_break()
    doc.add_heading("Evidence map and review", 1)
    intro = doc.add_paragraph("Proposed work for partner review. The record below does not approve a motion for filing.")
    intro.paragraph_format.space_after = Pt(10)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.25, 2.85, 2.9)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for index, title in enumerate(("Question", "Record and limit", "Proposed work")):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell_format(cell, "E8ECE9")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement("w:tblHeader"))
    for item in brief["evidence"]:
        cells = table.add_row().cells
        for index, key in enumerate(("question", "record", "action")):
            cells[index].width = Inches(widths[index])
            cell_format(cells[index])
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(item[key])
            run.font.size = Pt(10.5)
            if index == 0:
                run.bold = True
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_heading(brief["review"]["title"], 2)
    doc.add_paragraph(brief["review"]["body"])
    doc.add_heading("Decision for partner review", 2)
    doc.add_paragraph(brief["decision"])
    doc.add_heading("Source record", 2)
    for item in brief["sources"]:
        doc.add_paragraph(item["id"] + " · " + item["title"] + " · " + item["date"], "Source Reference")

    props = doc.core_properties
    props.title = brief["documentTitle"]
    props.subject = brief["name"]
    props.author = brief["author"]
    props.keywords = "Fictional demonstration; Northstar; internal draft; partner review"
    props.revision = brief["revision"]
    props.created = props.modified = datetime(2026, 10, 23, 12, 0, tzinfo=timezone.utc)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print("Created " + str(output))


if __name__ == "__main__":
    write(json.loads(sys.stdin.buffer.read().decode("utf-8-sig")), Path(sys.argv[1]))
