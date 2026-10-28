"""Build the internal Word reference from the same report model as the demo."""
import json
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'src/demo/assessment.json').read_text(encoding='utf-8'))
OUTPUT = ROOT / 'public' / DATA['downloadPath'].lstrip('/')

def paragraph(doc, text, style=None):
    return doc.add_paragraph(text, style)

def table(doc, headings, widths):
    result = doc.add_table(rows=1, cols=len(headings))
    result.autofit = False
    for cell, width, label in zip(result.rows[0].cells, widths, headings):
        cell.width = Inches(width)
        cell.text = label
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'EAF0F3')
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    result.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return result

def row(result, texts, widths):
    cells = result.add_row().cells
    for cell, text, width in zip(cells, texts, widths):
        cell.width = Inches(width)
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins = OxmlElement('w:tcMar')
        for edge in ['top', 'bottom', 'start', 'end']:
            item = OxmlElement('w:' + edge)
            item.set(qn('w:w'), '90')
            item.set(qn('w:type'), 'dxa')
            margins.append(item)
        cell._tc.get_or_add_tcPr().append(margins)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                run.font.size = Pt(10)
    result.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))

def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.75), Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance = section.footer_distance = Inches(.3)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(29)
    doc.styles['Title'].paragraph_format.space_after = Pt(6)
    doc.styles['Subtitle'].font.size = Pt(13)
    for name in ['Heading 1', 'Heading 2']:
        doc.styles[name].font.size = Pt(14 if name == 'Heading 1' else 11)
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.add_run().add_picture(str(ROOT / 'public/brand/asg/alliance-systems-group-mark-refined.png'), width=Inches(.27))
    header.add_run('  ALLIANCE SYSTEMS GROUP').font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Northstar Foods  |  Confidential  |  September 7, 2026  |  ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    paragraph(doc, DATA['title'], 'Title')
    paragraph(doc, 'Northstar Foods\nDecision rights and operating model', 'Subtitle')
    p = paragraph(doc, f"Prepared by {DATA['author']}  |  Reviewed by {DATA['reviewer']}\n{DATA['date']}  |  Approved revision {DATA['revision']}")
    for run in p.runs:
        run.font.size = Pt(9)
    doc.add_heading('Executive assessment', 1)
    paragraph(doc, DATA['decisionTitle']).runs[0].bold = True
    paragraph(doc, DATA['executiveSummary'])
    for decision in DATA['decisions']:
        doc.add_heading(decision['label'] + '  ' + decision['decision'], 2)
        paragraph(doc, decision['basis'])
    doc.add_heading('Assessment method', 1)
    for method in DATA['method']:
        p = paragraph(doc, '')
        p.add_run(method['criterion'] + '. ').bold = True
        p.add_run(method['test'])

    doc.add_page_break()
    doc.add_heading('Findings and decision limits', 1)
    for finding in DATA['findings']:
        doc.add_heading(finding['id'] + '  ' + finding['title'], 2)
        p = paragraph(doc, finding['source'])
        for run in p.runs:
            run.italic = True
            run.font.size = Pt(9)
        paragraph(doc, finding['observation'])
        for lead, content in [('Decision consequence. ', finding['consequence']), ('Release condition. ', finding['release'])]:
            p = paragraph(doc, '')
            p.add_run(lead).bold = True
            p.add_run(content)
    doc.add_heading('Assessment boundaries', 1)
    paragraph(doc, DATA['limitations'])

    doc.add_page_break()
    doc.add_heading('Actions before the diagnostic baseline', 1)
    paragraph(doc, 'Follow-up targets precede the September 11 baseline gate. Northstar owns its evidence; Avery coordinates validation and Sofia reviews the remaining conditions.')
    widths = [4.1, 1.9, 1.0]
    actions = table(doc, ['Action and completion condition', 'Accountable owner', 'Target'], widths)
    for action in DATA['actionPlan']:
        row(actions, [action['id'] + '  ' + action['action'] + '\n' + action['finding'] + '. ' + action['doneWhen'], action['owner'] + '\nCoordinate: ' + action['coordinator'], action['target']], widths)
    doc.add_heading('Reviewed revision', 1)
    paragraph(doc, DATA['reviewResponse'])
    doc.add_heading('Source register', 1)
    for source in DATA['sourceRegister']:
        p = paragraph(doc, '')
        p.add_run(source['id'] + '  ' + source['name'] + '. ').bold = True
        p.add_run(source['use'])
        for run in p.runs:
            run.font.size = Pt(10)
    props = doc.core_properties
    props.title, props.subject = DATA['title'], 'Northstar Foods operating-model diagnostic'
    props.author, props.last_modified_by = DATA['author'], DATA['reviewer']
    props.identifier, props.revision = 'NS-DR01-R2', DATA['revision']
    props.created = props.modified = datetime(2026, 9, 7, tzinfo=timezone.utc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    build()
